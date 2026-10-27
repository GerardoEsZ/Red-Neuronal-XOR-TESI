from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from pathlib import Path
base=Path(__file__).resolve().parent
D=Document();sec=D.sections[0];sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(3)
normal=D.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.space_after=Pt(4)
for name in ['Title','Heading 1','Heading 2']:
 s=D.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(23,42,69);s.font.bold=True;s.font.size=Pt(14);s.paragraph_format.space_before=Pt(6);s.paragraph_format.space_after=Pt(3)
def title(t):
 p=D.add_paragraph(style='Heading 1');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(t)
def p(txt):
 a=D.add_paragraph(txt);a.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;return a
def table(rows):
 tb=D.add_table(rows=0,cols=len(rows[0]));tb.style='Table Grid';tb.alignment=WD_TABLE_ALIGNMENT.CENTER
 for rr in rows:
  cells=tb.add_row().cells
  for cc,value in zip(cells,rr):cc.text=str(value)
 for cc in tb.rows[0].cells:
  for r in cc.paragraphs[0].runs:r.bold=True
 return tb
header=sec.header.paragraphs[0];header.text='TESI | Redes Neuronales Artificiales | Segundo parcial · Actividad 2';header.alignment=WD_ALIGN_PARAGRAPH.CENTER
header.runs[0].font.name='Arial';header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.text='Espitia Zavala Gerardo · Olvera Rodríguez Bruno | Grupo 1951';footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in footer.runs:r.font.size=Pt(9)
title('REPORTE DE RESULTADOS — ACTIVIDAD 2')
p('Del Perceptrón a la Retropropagación | Semestre 9.º · Grupo 1951 | Octubre de 2026')
title('1. Introducción y objetivo')
p('Se implementó en Python una red neuronal multicapa que aprende la compuerta XOR. La actividad permite observar directamente el paso hacia adelante, el cálculo de la pérdida, la retropropagación y la actualización de pesos, además de comparar la convergencia con diferentes tasas de aprendizaje.')
title('2. Metodología y herramientas')
p('Herramientas: Python, NumPy para álgebra matricial, Tkinter para la interfaz y Matplotlib para gráficas. Hardware: computadora personal compatible con Python. Se programó una red totalmente conectada 2–2–1 (dos entradas, dos neuronas ocultas y una salida), con activación sigmoide y sesgos entrenables. Se usó entropía cruzada binaria promedio y descenso por gradiente por lote completo.')
p('En cada época, la red ejecuta: (1) z₁ = XW₁ + b₁ y h = σ(z₁); (2) z₂ = hW₂ + b₂ y ŷ = σ(z₂); (3) pérdida BCE; (4) gradientes mediante regla de la cadena y (5) actualización W ← W − η∇W, b ← b − η∇b. Los datos son (0,0)→0, (0,1)→1, (1,0)→1 y (1,1)→0.')
p('Se realizaron dos entrenamientos desde exactamente los mismos pesos iniciales (semilla aleatoria 17), durante 10,000 épocas, modificando únicamente la tasa de aprendizaje: η=0.1 y η=1.0. El umbral de decisión fue 0.5.')
title('3. Resultados de los experimentos')
table([['Indicador','Experimento 1','Experimento 2'],['Tasa de aprendizaje','0.1','1.0'],['Épocas','10,000','10,000'],['Pérdida BCE final','0.0126158','0.0009507'],['Exactitud sobre 4 casos','100 %','100 %']])
p('Las probabilidades finales en el orden (0,0), (0,1), (1,0), (1,1) fueron: η=0.1 → [0.012254, 0.983709, 0.988766, 0.010357]; η=1.0 → [0.000910, 0.998746, 0.999154, 0.000790]. Todas las clases resultaron correctas al aplicar el umbral de 0.5.')
D.add_picture(str(base/'resultados/comparacion_loss.png'),width=Cm(14.7))
c=D.add_paragraph('Figura 1. Curvas de pérdida BCE por época con dos tasas de aprendizaje; resultados generados por el código Python.');c.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in c.runs:r.font.size=Pt(9)
title('4. Análisis de convergencia')
p('Ambos entrenamientos convergieron hacia soluciones que clasifican correctamente las cuatro entradas. Con η=1.0 la reducción de pérdida fue más rápida y el valor final menor bajo estas condiciones. Una tasa mayor no garantiza siempre mejores resultados: en otros problemas puede causar oscilaciones o inestabilidad. La arquitectura multicapa sí puede representar XOR, a diferencia del perceptrón simple.')
title('5. Sobreajuste y limitaciones')
p('No puede diagnosticarse sobreajuste (overfitting) de forma rigurosa con estos experimentos: el conjunto XOR contiene solamente cuatro casos y no se reservó un conjunto de prueba independiente. Tampoco se observa underfitting al final, ya que ambos modelos alcanzan las cuatro etiquetas correctamente. La exactitud del 100 % se refiere exclusivamente a estos cuatro patrones conocidos, no a generalización hacia nuevas tareas.')
title('6. Evidencia del funcionamiento y colaboración')
p('La pestaña «Laboratorio interactivo» muestra entradas, pesos y sesgos, activaciones ocultas, predicción y gradientes de la última época. La pestaña «Experimentos y gráficas» permite repetir las dos pruebas y exportar los datos a CSV y la curva a PNG. El archivo red_xor.py contiene las ecuaciones de aprendizaje y simulador_xor.py controla la visualización. La evidencia de participación de cada integrante debe completarse con la distribución real de actividades y la validación de ambos.')
title('7. Conclusiones individuales (para revisión)')
p('Espitia Zavala Gerardo: la implementación de XOR evidencia que la capa oculta resuelve una limitación de representación del perceptrón simple. Inspeccionar las activaciones y los gradientes permite comprender por qué disminuye la pérdida durante el entrenamiento.')
p('Olvera Rodríguez Bruno: comparar dos tasas de aprendizaje con idéntica inicialización demuestra cómo un hiperparámetro influye en la velocidad de convergencia. También es importante distinguir la memorización de cuatro patrones de la capacidad de generalización.')
title('8. Referencias y apoyo tecnológico')
p('Rumelhart, D. E., Hinton, G. E., & Williams, R. J. (1986). Learning representations by back-propagating errors. Nature, 323, 533–536. https://doi.org/10.1038/323533a0')
p('OpenAI. (2026). ChatGPT [Modelo de lenguaje]. https://chatgpt.com/ (Apoyo en documentación y preparación de materiales; los resultados proceden de ejecuciones del programa).')
path=base/'Reporte_Actividad_2.docx';D.save(path);print(path)
